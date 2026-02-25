import streamlit as st
import sqlite3
import bcrypt
from pypdf import PdfReader
from docx import Document
import openpyxl

st.set_page_config(page_title="Buscador de Documentos", layout="wide")

# =========================
# BANCO
# =========================
def inicializar_banco():
    conn = sqlite3.connect("documentos.db")
    cursor = conn.cursor()

    cursor.execute("""
    CREATE TABLE IF NOT EXISTS usuarios (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        username TEXT UNIQUE,
        senha_hash TEXT
    )
    """)

    cursor.execute("""
    CREATE TABLE IF NOT EXISTS documentos (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        nome TEXT,
        arquivo BLOB
    )
    """)

    cursor.execute("""
    CREATE VIRTUAL TABLE IF NOT EXISTS documentos_fts
    USING fts5(nome, conteudo)
    """)

    cursor.execute("SELECT * FROM usuarios WHERE username = ?", ("admin",))
    if not cursor.fetchone():
        senha_hash = bcrypt.hashpw("admin123".encode(), bcrypt.gensalt()).decode()
        cursor.execute(
            "INSERT INTO usuarios (username, senha_hash) VALUES (?, ?)",
            ("admin", senha_hash),
        )

    conn.commit()
    conn.close()


inicializar_banco()


# =========================
# FUNÇÕES
# =========================
def verificar_login(username, senha):
    conn = sqlite3.connect("documentos.db")
    cursor = conn.cursor()

    cursor.execute("SELECT senha_hash FROM usuarios WHERE username = ?", (username,))
    resultado = cursor.fetchone()
    conn.close()

    if resultado:
        return bcrypt.checkpw(senha.encode(), resultado[0].encode())
    return False


def extrair_texto(arquivo):
    nome = arquivo.name
    extensao = nome.split(".")[-1].lower()

    if extensao == "pdf":
        reader = PdfReader(arquivo)
        texto = ""
        for page in reader.pages:
            texto += page.extract_text() or ""
        return texto

    elif extensao == "docx":
        doc = Document(arquivo)
        return "\n".join([p.text for p in doc.paragraphs])

    elif extensao == "xlsx":
        wb = openpyxl.load_workbook(arquivo)
        texto = ""
        for sheet in wb:
            for row in sheet.iter_rows(values_only=True):
                texto += " ".join([str(cell) for cell in row if cell]) + "\n"
        return texto

    elif extensao == "txt":
        return arquivo.read().decode("utf-8")

    return ""


def salvar_documento(nome, conteudo, arquivo_bytes):
    conn = sqlite3.connect("documentos.db")
    cursor = conn.cursor()

    cursor.execute(
        "INSERT INTO documentos (nome, arquivo) VALUES (?, ?)",
        (nome, arquivo_bytes),
    )

    cursor.execute(
        "INSERT INTO documentos_fts (nome, conteudo) VALUES (?, ?)",
        (nome, conteudo),
    )

    conn.commit()
    conn.close()


# =========================
# LOGIN
# =========================
if "logado" not in st.session_state:
    st.session_state.logado = False

if not st.session_state.logado:

    st.title("🔐 Login")

    username = st.text_input("Usuário")
    senha = st.text_input("Senha", type="password")

    if st.button("Entrar"):
        if verificar_login(username, senha):
            st.session_state.logado = True
            st.rerun()
        else:
            st.error("Usuário ou senha inválidos.")

else:
    st.title("📂 Buscador de Documentos")

    # =========================
    # UPLOAD
    # =========================
    st.subheader("📤 Upload de Documento")

    arquivo = st.file_uploader(
        "Envie um arquivo",
        type=["pdf", "docx", "xlsx", "txt"]
    )

    if arquivo:
        with st.spinner("Processando e indexando..."):
            texto = extrair_texto(arquivo)
            if texto.strip():
                salvar_documento(arquivo.name, texto, arquivo.getvalue())
                st.success("Documento salvo e indexado com sucesso!")
            else:
                st.warning("Não foi possível extrair texto.")

    # =========================
    # BUSCA
    # =========================
    st.subheader("🔎 Buscar Documento")

    busca = st.text_input("Digite o termo")

    if busca:
        conn = sqlite3.connect("documentos.db")
        cursor = conn.cursor()

        cursor.execute("""
            SELECT d.id, d.nome, d.arquivo
            FROM documentos d
            JOIN documentos_fts f ON d.nome = f.nome
            WHERE documentos_fts MATCH ?
        """, (busca,))

        resultados = cursor.fetchall()
        conn.close()

        if resultados:
            for doc_id, nome, arquivo_blob in resultados:
                st.write("📄", nome)

                st.download_button(
                    label="⬇️ Baixar",
                    data=arquivo_blob,
                    file_name=nome,
                    key=doc_id
                )
        else:
            st.warning("Nenhum resultado encontrado.")

    if st.button("Logout"):
        st.session_state.logado = False
        st.rerun()
