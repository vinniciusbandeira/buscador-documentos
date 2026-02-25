import os
import sqlite3
from pypdf import PdfReader
from docx import Document
import openpyxl

PASTA_DOCUMENTOS = r"documentos"

conn = sqlite3.connect("documentos.db")
cursor = conn.cursor()

cursor.execute("""
CREATE VIRTUAL TABLE IF NOT EXISTS documentos_fts
USING fts5(nome, caminho, conteudo)
""")

def extrair_texto(caminho):
    ext = os.path.splitext(caminho)[1].lower()

    try:
        if ext == ".txt":
            with open(caminho, "r", encoding="utf-8", errors="ignore") as f:
                return f.read()

        elif ext == ".pdf":
            reader = PdfReader(caminho)
            texto = ""
            for page in reader.pages:
                texto += page.extract_text() or ""
            return texto

        elif ext == ".docx":
            doc = Document(caminho)
            return "\n".join([p.text for p in doc.paragraphs])

        elif ext == ".xlsx":
            wb = openpyxl.load_workbook(caminho)
            texto = ""
            for sheet in wb:
                for row in sheet.iter_rows(values_only=True):
                    texto += " ".join([str(cell) for cell in row if cell]) + "\n"
            return texto

    except:
        return ""

    return ""

for raiz, dirs, arquivos in os.walk(PASTA_DOCUMENTOS):
    for arquivo in arquivos:
        caminho_completo = os.path.join(raiz, arquivo)

        texto = extrair_texto(caminho_completo)

        if texto.strip():
            cursor.execute(
                "INSERT INTO documentos_fts (nome, caminho, conteudo) VALUES (?, ?, ?)",
                (arquivo, caminho_completo, texto)
            )
            print("Indexado:", caminho_completo)

conn.commit()
conn.close()

print("\nBanco criado e documentos indexados com sucesso!")